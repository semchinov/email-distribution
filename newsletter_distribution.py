#!/usr/bin/env python3
"""Send newsletter emails to recipients listed in an Excel file."""

from __future__ import annotations

import argparse
import smtplib
import ssl
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Distribute an email newsletter using recipients from an Excel file.",
    )
    parser.add_argument("excel_file", type=Path, help="Path to recipients .xlsx file")
    parser.add_argument("subject", help="Email subject")
    parser.add_argument("body_file", type=Path, help="Path to plain-text newsletter body")
    parser.add_argument("--smtp-host", required=True, help="SMTP host")
    parser.add_argument("--smtp-port", type=int, default=587, help="SMTP port (default: 587)")
    parser.add_argument("--smtp-user", required=True, help="SMTP username")
    parser.add_argument("--smtp-password", required=True, help="SMTP password")
    parser.add_argument("--sender", required=True, help="From email address")
    parser.add_argument("--sheet", help="Excel sheet name (default: active sheet)")
    parser.add_argument(
        "--email-column",
        default="email",
        help="Column header with recipient emails (default: email)",
    )
    parser.add_argument(
        "--name-column",
        default="name",
        help="Column header with recipient names (default: name)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Do not send emails, only print who would receive them",
    )
    return parser.parse_args()


def read_recipients(excel_file: Path, sheet: str | None, email_column: str, name_column: str) -> list[dict[str, str]]:
    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active

    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError("Excel file does not have a header row.")

    normalized_headers = {
        str(header).strip().lower(): idx
        for idx, header in enumerate(headers)
        if header is not None and str(header).strip()
    }

    email_idx = normalized_headers.get(email_column.strip().lower())
    if email_idx is None:
        raise ValueError(f"Required email column '{email_column}' not found.")

    name_idx = normalized_headers.get(name_column.strip().lower())

    recipients: list[dict[str, str]] = []
    for row in rows:
        if row is None:
            continue

        email_value = row[email_idx] if email_idx < len(row) else None
        if not email_value:
            continue

        email = str(email_value).strip()
        if "@" not in email:
            continue

        name = ""
        if name_idx is not None and name_idx < len(row) and row[name_idx] is not None:
            name = str(row[name_idx]).strip()

        recipients.append({"email": email, "name": name})

    return recipients


def build_message(sender: str, recipient: dict[str, str], subject: str, template: str) -> EmailMessage:
    message = EmailMessage()
    message["From"] = sender
    message["To"] = recipient["email"]
    message["Subject"] = subject
    body = template.replace("{name}", recipient["name"]).replace("{email}", recipient["email"])
    message.set_content(body)
    return message


def send_newsletter(args: argparse.Namespace) -> int:
    if not args.excel_file.exists():
        raise FileNotFoundError(f"Excel file not found: {args.excel_file}")
    if not args.body_file.exists():
        raise FileNotFoundError(f"Body file not found: {args.body_file}")

    recipients = read_recipients(args.excel_file, args.sheet, args.email_column, args.name_column)
    if not recipients:
        raise ValueError("No valid recipients found in the Excel file.")

    body_template = args.body_file.read_text(encoding="utf-8")

    if args.dry_run:
        for recipient in recipients:
            print(f"[DRY RUN] Would send to {recipient['email']}")
        print(f"Prepared {len(recipients)} emails.")
        return 0

    with smtplib.SMTP(args.smtp_host, args.smtp_port, timeout=30) as smtp:
        smtp.starttls(context=ssl.create_default_context())
        smtp.login(args.smtp_user, args.smtp_password)
        for recipient in recipients:
            smtp.send_message(build_message(args.sender, recipient, args.subject, body_template))
            print(f"Sent to {recipient['email']}")

    print(f"Successfully sent {len(recipients)} emails.")
    return 0


def main() -> int:
    args = parse_args()
    return send_newsletter(args)


if __name__ == "__main__":
    raise SystemExit(main())
